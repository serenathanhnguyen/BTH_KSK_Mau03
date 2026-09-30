# -*- coding: utf-8 -*-
"""Kiem tra nhanh (khong dung pytest de don gian): chay het pipeline tren
sample_data/mau_03_sample_fake.xlsx (du lieu hu cau) va in ket qua."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import json
import reader
import mapping as mapping_mod
import pipeline
import report as report_mod

def main():
    rr = reader.read_m03("sample_data/mau_03_sample_fake.xlsx")
    assert rr.total_valid_rows == 3, f"Expected 3 people, got {rr.total_valid_rows}"
    print("OK reader:", rr.total_valid_rows, "nguoi, keyword row =", rr.keyword_row)

    cfg = mapping_mod.load_mapping("mapping.json")
    with open("lab_reference.json", encoding="utf-8") as f:
        lab_ref = json.load(f)
    with open("cbc_reference.json", encoding="utf-8") as f:
        cbc_ref = json.load(f)

    manual_map = {"alt": "kskdk_shm_alat_gpt", "ast": "kskdk_shm_asat_got",
                  "ure": "kskdk_shm_ure", "creatinin": "kskdk_shm_creatinin",
                  "glucose": "kskdk_shm_duongmau",
                  "acid_uric": None, "cholesterol": None, "triglycerid": None,
                  "hdl": None, "ldl": None, "satq": None}

    records = pipeline.process_all(rr, manual_map, lab_ref, cbc_ref,
                                    hb_source_unit="g/dL", urine_enabled=True,
                                    icd_fallback_enabled=True, sort_by_ten=True)
    assert len(records) == 3
    print("OK pipeline: 3 ban ghi")

    for r in records:
        print(" -", r["stt"], r["ten"], "| Xếp loại:", r["xeploai"], "| CTM:", r["ctm"],
              "| Ghi chú:", r["ghichu"], "| Cảnh báo:", r["canhbao"][:80])

    # kiem tra thi luc: nguoi 2 (Trần Thị B) chi co cap kinh lo (8,9) -> phai chon dung cap nay
    b = [r for r in records if r["cccd"] == "000000000002"][0]
    assert (b["matphai"], b["mattrai"]) == (8, 9), b
    print("OK vision pair chon dung cap co du lieu")

    # kiem tra CTM: nguoi 2 co Hb=10.5 (<11.0 -> giam), WBC=11.2 (>10.0 -> tang)
    assert "Hb giảm" in b["ctm"] and "tăng" in b["ctm"], b["ctm"]
    print("OK CTM ghep dung thu tu Hb->MCV->WBC->PLT:", b["ctm"])

    avail = pipeline.column_availability(records)
    enabled_cols = mapping_mod.ordered_enabled_columns(cfg, availability=avail)
    enabled_ids = [c["id"] for c in enabled_cols]
    assert "cho" not in enabled_ids and "tri" not in enabled_ids, "Cot Cholesterol/Triglycerid phai bi an vi khong co du lieu"
    assert "glu" in enabled_ids, "Cot Glucose phai HIEN vi sample co du lieu glucose (kiem tra manual_map co map dung keyword)"
    for r in records:
        assert r["lab"]["glucose"]["value"] is not None, f"Glucose bi None cho {r['ten']} - kiem tra keyword mapping"
    print("OK auto-hide cot khong co du lieu:", enabled_ids)

    meta = dict(don_vi_chu_quan="ỦY BAN NHÂN DÂN PHƯỜNG X", tram_y_te="TRẠM Y TẾ",
                doi_tuong="CB-NV TEST", nam="2026", nguoi_lap_bang="Test", giam_doc="")
    out_path = "/tmp/smoke_test_output.xlsx"
    report_mod.build_report(records, enabled_cols, meta, out_path)
    assert os.path.exists(out_path)
    print("OK report.build_report ->", out_path)

    # doc lai file xuat ra de kiem tra co du 3 dong + tieu de
    import openpyxl
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    assert ws["A9"].value == 1
    assert ws["A11"].value == 3
    print("OK file xuat co du 3 dong du lieu (row 9-11)")
    print("\nTAT CA KIEM TRA DEU PASS.")

if __name__ == "__main__":
    main()
