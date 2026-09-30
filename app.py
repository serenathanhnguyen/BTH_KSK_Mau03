"""Ứng dụng khung tổng hợp KSK M03 → mẫu BTH/MEO.
Đổi tên file này thành app.py trên GitHub; thêm requirements.txt với:
streamlit>=1.35,<2
openpyxl>=3.1,<4
"""
import io
import json
import math
from copy import copy
from datetime import date, datetime

import openpyxl
import streamlit as st
from openpyxl.utils import get_column_letter


st.set_page_config(page_title="Tổng hợp KSK", layout="wide")
st.title("Bảng tổng hợp khám sức khỏe")
st.caption("M03: keyword ở dòng 4, dữ liệu từ dòng 5. BTH: mã cột ở dòng 6, danh sách từ dòng 9.")

# Giá trị là keyword dòng 4 của M03. None = chờ xác nhận nguồn dữ liệu.
# Nội khoa M03 có nhiều phân hệ nên không tự gộp thành một phân loại.
DEFAULT = {
    "hoten": "ho_ten", "cccd": "dinh_danh_ca_nhan",
    "ngaysinh": "ngay_sinh", "gioitinh": "gioi_tinh",
    "chieucao": "chieucao", "cannang": "cannang",
    "ngoai": "ngoaikhoa_phanloai", "dalieu": "dalieu_phanloai",
    "sanphukhoa": "phukhoa_phanloai", "mat": "mat_phanloai",
    "taimuihong": "tmh_phanloai", "ranghammat": "rhm_phanloai",
    "alt": "kskdk_shm_alat_gpt", "ast": "kskdk_shm_asat_got",
    "ure": "kskdk_shm_ure", "cre": "kskdk_shm_creatinin",
    "glu": "kskdk_shm_duongmau",
}
SPECIAL = {"stt", "ten", "bmi", "huyetap", "matphai", "mattrai"}
MEO_COLUMNS = ["stt", "hoten", "ten", "cccd", "ngaysinh", "gioitinh",
               "chieucao", "cannang", "bmi", "huyetap", "matphai",
               "mattrai", "noi", "ngoai", "dalieu", "sanphukhoa",
               "mat", "taimuihong", "ranghammat", "xeploai",
               "ghichu", "canhbao"]


def display(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    return str(value).strip()


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = display(value)
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    return value  # Không đoán ngày nếu dữ liệu không đúng định dạng.


def source_schema(ws):
    result = {}
    for col in range(1, ws.max_column + 1):
        keyword = display(ws.cell(4, col).value)
        if keyword:
            # File M03 có keyword mat_docau_mt bị lặp; dùng cột đầu tiên.
            result.setdefault(keyword, col)
    return result


def output_schema(ws):
    result = {}
    for col in range(1, 36):
        keyword = display(ws.cell(6, col).value)
        if keyword:
            result[keyword] = col
    if not result and ws["A7"].value == "STT" and ws["T7"].value == "Xếp loại":
        return {key: index for index, key in enumerate(MEO_COLUMNS, 1)}
    return result


def source_rows(ws, schema):
    identity = schema.get("dinh_danh_ca_nhan")
    name = schema.get("ho_ten")
    if not identity or not name:
        raise ValueError("Thiếu keyword dinh_danh_ca_nhan hoặc ho_ten trong file dữ liệu.")
    for row in ws.iter_rows(min_row=5, max_col=ws.max_column, values_only=True):
        if display(row[name - 1]) or display(row[identity - 1]):
            yield row


def take(row, schema, keyword):
    col = schema.get(keyword)
    return row[col - 1] if col and col <= len(row) else None


def convert(key, value):
    if key == "cccd":
        return display(value)  # Giữ số 0 đầu: nguồn phải lưu CCCD dạng text.
    if key == "ngaysinh":
        return parse_date(value)
    if key == "gioitinh":
        return {"1": "Nam", "2": "Nữ", "3": "Chưa xác định"}.get(display(value), display(value))
    return value


def calculated(key, row, schema, position, acuity):
    if key == "stt":
        return position
    if key == "ten":
        name = display(take(row, schema, "ho_ten"))
        return name.split()[-1] if name else None
    if key == "huyetap":
        systolic = display(take(row, schema, "huyetaptamthu"))
        diastolic = display(take(row, schema, "huyetaptamtruong"))
        return f"{systolic}/{diastolic}" if systolic and diastolic else None
    if key in ("matphai", "mattrai"):
        suffix = "mp" if key == "matphai" else "mt"
        return take(row, schema, f"mat_{acuity}_{suffix}")
    if key == "bmi":
        try:
            height = float(take(row, schema, "chieucao"))
            weight = float(take(row, schema, "cannang"))
            return round(weight / ((height / 100) ** 2), 2) if height > 0 else None
        except (TypeError, ValueError, ZeroDivisionError):
            return None
    return None


def make_report(template_bytes, records, schema, mapping, active, acuity, title, subtitle):
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    ws = wb.active
    columns = output_schema(ws)
    is_meo = columns.get("xeploai") == 20 and not ws["A6"].value
    end_col = max(columns.values())
    class_col = get_column_letter(columns["xeploai"])
    extra = max(0, len(records) - 10)
    if extra:
        # openpyxl không tự chuyển merged ranges khi chèn dòng.
        original_merges = [str(r) for r in ws.merged_cells.ranges]
        for region in original_merges:
            ws.unmerge_cells(region)
        ws.insert_rows(19, extra)
        for region in original_merges:
            bounds = openpyxl.utils.range_boundaries(region)
            c1, r1, c2, r2 = bounds
            if r1 >= 19:
                r1 += extra
                r2 += extra
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        for target_row in range(19, 19 + extra):
            ws.row_dimensions[target_row].height = ws.row_dimensions[9].height
            for col in range(1, end_col + 1):
                src, dst = ws.cell(9, col), ws.cell(target_row, col)
                dst._style = copy(src._style)
                dst.alignment = copy(src.alignment)
                dst.protection = copy(src.protection)

    ws["A4"] = title
    ws["A5"] = subtitle
    # Xóa 10 dòng dữ liệu mẫu cũ và các dòng mở rộng, giữ phần ký tên.
    for row_idx in range(9, 19 + extra):
        for col in range(1, end_col + 1):
            ws.cell(row_idx, col).value = None

    for i, record in enumerate(records, 1):
        row_idx = 8 + i
        for key, col in columns.items():
            if key not in active:
                continue
            value = (calculated(key, record, schema, i, acuity) if key in SPECIAL
                     else take(record, schema, mapping.get(key)))
            value = convert(key, value)
            cell = ws.cell(row_idx, col)
            cell.value = value if value != "" else None
            if key == "cccd":
                cell.number_format = "@"
            elif key == "ngaysinh" and isinstance(value, date):
                cell.number_format = "dd/mm/yyyy"
        note_col = columns.get("ghichu")
        if note_col:
            note = display(ws.cell(row_idx, note_col).value)
            if note:
                letter = get_column_letter(note_col)
                width = ws.column_dimensions[letter].width or 20
                lines = max(1, math.ceil(len(note) / max(8, width - 3)))
                minimum = ws.row_dimensions[row_idx].height or ws.row_dimensions[9].height or 18.85
                ws.row_dimensions[row_idx].height = max(minimum, lines * 12 + 4)

    first, last = 9, max(9, 8 + len(records))
    total = 19 + extra
    # Giữ nguyên phông chữ, đường viền, căn lề và ô gộp của mẫu.
    ws.cell(total, 1).value = "TỔNG CỘNG"
    for col in range(3, end_col):
        letter = get_column_letter(col)
        cell = ws.cell(total, col)
        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
            cell.value = f'=COUNTA({letter}{first}:{letter}{last})' if records else 0
    for index, report_row in enumerate(range(total + 3, total + 8), 1):
        ws.cell(report_row, 3).value = (
            f'=COUNTIF(${class_col}${first}:${class_col}${last},{index})' if records else 0
        )
    ws.cell(total + 2, 3).value = len(records)
    ws.print_area = f"A1:{'U' if is_meo else 'AI'}{35 + extra}"
    ws.print_title_rows = "7:8"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.freeze_panes = "G9" if is_meo else ws.freeze_panes
    # V của mẫu MEO là cảnh báo nội bộ, phải luôn ẩn và không in.
    if is_meo:
        ws.column_dimensions["V"].hidden = True
    else:
        for key, col in columns.items():
            ws.column_dimensions[get_column_letter(col)].hidden = key not in active
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


data_file = st.file_uploader("1. File dữ liệu M03 (.xlsx)", type="xlsx", key="data")
template_file = st.file_uploader("2. File mẫu báo cáo BTH (.xlsx)", type="xlsx", key="template")
if data_file and template_file:
    try:
        data_book = openpyxl.load_workbook(io.BytesIO(data_file.getvalue()), read_only=True, data_only=True)
        data_sheet = st.selectbox("Sheet dữ liệu", data_book.sheetnames,
                                  index=data_book.sheetnames.index("ThongTinHanhChinh")
                                  if "ThongTinHanhChinh" in data_book.sheetnames else 0)
        data_ws = data_book[data_sheet]
        schema = source_schema(data_ws)
        records = list(source_rows(data_ws, schema))
        template_book = openpyxl.load_workbook(io.BytesIO(template_file.getvalue()), read_only=True)
        columns = output_schema(template_book.active)
        if not columns or "hoten" not in columns:
            raise ValueError("Không nhận diện được mẫu BTH/MEO: cần keyword dòng 6 hoặc tiêu đề MEO ở dòng 7.")
        is_meo = columns.get("xeploai") == 20 and not template_book.active["A6"].value
        if is_meo:
            st.caption("Mẫu MEO: A–U là báo cáo; V là cột cảnh báo ẩn. Định dạng, độ rộng, ô gộp và ký tên lấy từ file mẫu.")
        st.info(f"Đọc được {len(records):,} dòng có họ tên hoặc CCCD. Hãy kiểm tra dòng trống và bản ghi trùng trước khi xuất.")
        with st.expander("3. Ánh xạ cột và chọn cột xuất", expanded=True):
            keys = list(schema)
            active = (list(columns) if is_meo else st.multiselect(
                "Cột hiển thị trong báo cáo", list(columns), default=list(columns)))
            acuity = st.selectbox("Thị lực P/T lấy từ", ["khongkinh", "cokinh", "kinhlo"],
                                  format_func=lambda x: {"khongkinh":"Không kính", "cokinh":"Có kính", "kinhlo":"Kính lỗ"}[x])
            mapping = {}
            for key in columns:
                if key in SPECIAL or key not in active:
                    continue
                choices = ["(Để trống)"] + keys
                suggested = DEFAULT.get(key)
                selection = st.selectbox(
                    f"{get_column_letter(columns[key])} · {key}", choices,
                    index=choices.index(suggested) if suggested in choices else 0,
                    key=f"map_{key}",
                )
                mapping[key] = None if selection == "(Để trống)" else selection
        st.caption("Xếp loại T/AH chưa có keyword tổng thể được xác nhận; chọn thủ công khi bạn xác định đúng cột. Nội khoa M gồm nhiều phân hệ nên mặc định để trống.")
        title = st.text_input("Tên bảng", "BẢNG TỔNG HỢP PHÂN LOẠI SỨC KHỎE")
        subtitle = st.text_input("Dòng năm / đợt khám", "KHÁM SỨC KHỎE NĂM 2026")
        st.download_button("Tải cấu hình ánh xạ JSON", json.dumps({
            "mapping": mapping, "active": active, "acuity": acuity
        }, ensure_ascii=False, indent=2).encode("utf-8"), "anh_xa_bao_cao_ksk.json", "application/json")
        if st.button("Tạo báo cáo", type="primary"):
            result = make_report(template_file.getvalue(), records, schema, mapping, active,
                                 acuity, title, subtitle)
            st.download_button("Tải bảng tổng hợp KSK", result, "BTH_KSK_tong_hop.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as exc:
        st.error(f"Không tạo được báo cáo: {exc}")
