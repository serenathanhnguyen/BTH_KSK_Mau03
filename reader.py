# -*- coding: utf-8 -*-
"""
reader.py
Doc file M03 (sheet ThongTinHanhChinh) theo KEYWORD o hang 4, khong dua vao
vi tri cot co dinh (muc 1 va muc 4.1 cua khung nguyen tac - vi tri cot co
the doi giua cac ban xuat khac nhau cua Medinet).
"""
from openpyxl import load_workbook

SHEET_NAME = "ThongTinHanhChinh"
KEYWORD_ROW_CANDIDATES = (4, 3, 5)  # quet linh hoat quanh hang 4 phong khi mau lech dong
DATA_START_ROW_DEFAULT = 5
BLANK_STREAK_STOP = 6  # so dong trong lien tiep de coi la da het du lieu (muc "doc den dong cuoi cung thuc su")

ANCHOR_KEYWORDS = ("dinh_danh_ca_nhan", "ho_ten")


class ReaderResult:
    def __init__(self):
        self.keyword_row = None
        self.kw_to_cols = {}          # keyword -> list[col_idx]  (list vi co the trung, vd mat_docau_mt)
        self.duplicate_keywords = []  # cac keyword xuat hien >1 lan
        self.rows = []                # list[dict keyword -> value] tho, 1 dict / nguoi
        self.total_rows_scanned = 0
        self.total_valid_rows = 0
        self.skipped_rows = []        # list[(row_idx, ly_do)]
        self.data_start_row = DATA_START_ROW_DEFAULT
        self.data_end_row = None


def _find_keyword_row(ws):
    for r in KEYWORD_ROW_CANDIDATES:
        found = 0
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip() in ANCHOR_KEYWORDS:
                found += 1
        if found >= 1:
            return r
    return None


def read_m03(file_path_or_buffer):
    wb = load_workbook(file_path_or_buffer, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"Không tìm thấy sheet '{SHEET_NAME}' trong file. Các sheet có sẵn: {wb.sheetnames}"
        )
    ws = wb[SHEET_NAME]

    result = ReaderResult()
    kw_row = _find_keyword_row(ws)
    if kw_row is None:
        raise ValueError(
            "Không tìm thấy hàng keyword (hàng 4) — kiểm tra lại đúng định dạng file Mẫu 03."
        )
    result.keyword_row = kw_row
    result.data_start_row = kw_row + 1

    kw_to_cols = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=kw_row, column=c).value
        if isinstance(v, str) and v.strip():
            kw_to_cols.setdefault(v.strip(), []).append(c)
    result.kw_to_cols = kw_to_cols
    result.duplicate_keywords = [k for k, cols in kw_to_cols.items() if len(cols) > 1]

    def first_col(keyword):
        cols = kw_to_cols.get(keyword)
        return cols[0] if cols else None

    r = result.data_start_row
    blank_streak = 0
    last_valid_row = None
    while r <= ws.max_row:
        anchor_cccd = ws.cell(row=r, column=first_col("dinh_danh_ca_nhan")).value if first_col("dinh_danh_ca_nhan") else None
        anchor_name = ws.cell(row=r, column=first_col("ho_ten")).value if first_col("ho_ten") else None
        result.total_rows_scanned += 1

        if (anchor_cccd in (None, "")) and (anchor_name in (None, "")):
            blank_streak += 1
            if blank_streak >= BLANK_STREAK_STOP:
                break
            r += 1
            continue
        blank_streak = 0

        if anchor_name in (None, ""):
            result.skipped_rows.append((r, "Thiếu họ tên"))
            r += 1
            continue

        row_data = {"_excel_row": r}
        for kw, cols in kw_to_cols.items():
            # neu trung keyword, luu ca list gia tri de UI cho nguoi dung chon
            vals = [ws.cell(row=r, column=c).value for c in cols]
            row_data[kw] = vals[0] if len(vals) == 1 else vals
        result.rows.append(row_data)
        result.total_valid_rows += 1
        last_valid_row = r
        r += 1

    result.data_end_row = last_valid_row
    return result
